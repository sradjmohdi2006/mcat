from cat_tool.formats._text_base import make_segments_from_texts


class PptxHandler:
    def load(self, file_path):
        from pptx import Presentation
        prs = Presentation(file_path)
        texts = []
        locations = []
        for slide_num, slide in enumerate(prs.slides):
            for shape_idx, shape in enumerate(slide.shapes):
                if shape.has_text_frame:
                    for para_idx, para in enumerate(shape.text_frame.paragraphs):
                        t = para.text.strip()
                        if t:
                            texts.append(t)
                            locations.append({"slide": slide_num, "shape": shape_idx, "paragraph": para_idx})
                if shape.has_table:
                    for ri, row in enumerate(shape.table.rows):
                        for ci, cell in enumerate(row.cells):
                            t = cell.text.strip()
                            if t:
                                texts.append(t)
                                locations.append({"slide": slide_num, "shape": shape_idx, "table_row": ri, "table_col": ci})
        return make_segments_from_texts(texts, "Slide ", source_locations=locations)

    def render(self, original_path, segments, output_path):
        from pptx import Presentation
        prs = Presentation(original_path)
        for seg in segments:
            loc = seg.get("source_location", {})
            if not loc:
                continue
            target = seg.get("target", "").strip()
            if not target:
                continue
            slide = prs.slides[loc["slide"]]
            shape = slide.shapes[loc["shape"]]
            if "paragraph" in loc:
                shape.text_frame.paragraphs[loc["paragraph"]].text = target
            elif "table_row" in loc:
                shape.table.cell(loc["table_row"], loc["table_col"]).text = target
        prs.save(output_path)
