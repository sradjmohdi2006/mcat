from cat_tool.formats._text_base import make_segments_from_texts


class XlsxHandler:
    def load(self, file_path):
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        texts = []
        locations = []
        for si, ws in enumerate(wb.worksheets):
            for ri, row in enumerate(ws.iter_rows(values_only=True)):
                for ci, cell in enumerate(row):
                    if cell is not None:
                        t = str(cell).strip()
                        if t:
                            texts.append(t)
                            locations.append({"sheet": si, "row": ri, "col": ci})
        wb.close()
        return make_segments_from_texts(texts, "Cell ", source_locations=locations)

    def render(self, original_path, segments, output_path):
        from openpyxl import load_workbook
        wb = load_workbook(original_path)
        for seg in segments:
            loc = seg.get("source_location", {})
            if not loc:
                continue
            target = seg.get("target", "").strip()
            if not target:
                continue
            ws = wb.worksheets[loc["sheet"]]
            ws.cell(row=loc["row"] + 1, column=loc["col"] + 1, value=target)
        wb.save(output_path)
