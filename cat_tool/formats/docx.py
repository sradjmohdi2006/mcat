from cat_tool.formats._text_base import make_segments_from_texts


class DocxHandler:
    def load(self, file_path):
        from docx import Document
        doc = Document(file_path)
        texts = []
        locations = []
        for i, p in enumerate(doc.paragraphs):
            t = p.text.strip()
            if t:
                texts.append(t)
                locations.append({"type": "paragraph", "index": i})
        for ti, table in enumerate(doc.tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    t = cell.text.strip()
                    if t:
                        texts.append(t)
                        locations.append({"type": "table_cell", "table_index": ti, "row_index": ri, "col_index": ci})
        return make_segments_from_texts(texts, "Para ", source_locations=locations)

    def render(self, original_path, segments, output_path):
        from docx import Document
        doc = Document(original_path)
        seg_iter = iter(segments)
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                seg = next(seg_iter, None)
                if seg and seg.get("source_location", {}).get("type") == "paragraph":
                    target = seg.get("target", "").strip()
                    if target:
                        p.clear()
                        p.add_run(target)
        for ti, table in enumerate(doc.tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    t = cell.text.strip()
                    if t:
                        seg = next(seg_iter, None)
                        if seg and seg.get("source_location", {}).get("type") == "table_cell":
                            target = seg.get("target", "").strip()
                            if target:
                                cell.text = target
        doc.save(output_path)

    def load_bookmarks(self, file_path):
        """Extract heading paragraphs as bookmarks from a DOCX file."""
        from docx import Document
        bookmarks = []
        try:
            doc = Document(file_path)
            heading_levels = {"Heading 1": 0, "Heading 2": 1, "Heading 3": 2,
                              "Heading 4": 3, "Heading 5": 4, "Heading 6": 5}
            for i, para in enumerate(doc.paragraphs):
                style_name = para.style.name if para.style else ""
                if style_name.startswith("Heading"):
                    level = heading_levels.get(style_name, 0)
                    title = para.text.strip()
                    if title:
                        bookmarks.append({
                            "title": title,
                            "level": level,
                            "page": None,
                            "para_index": i,
                        })
        except Exception:
            pass
        return bookmarks
