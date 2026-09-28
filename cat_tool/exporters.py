import os


def export_docx(segments, output_path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    doc.add_heading("Translation Export", level=1)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    hdr[0].text = "#"
    hdr[1].text = "Source"
    hdr[2].text = "Target"

    for idx, seg in enumerate(segments):
        row = table.add_row().cells
        row[0].text = str(idx + 1)
        row[1].text = seg.get("source_clean", seg["source"])
        tgt = seg.get("target_clean", seg["target"])
        row[2].text = tgt if tgt.strip() else ""

    doc.save(output_path)


def export_xlsx(segments, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Translations"
    ws.append(["#", "Source", "Target", "Status"])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 12

    for idx, seg in enumerate(segments):
        src = seg.get("source_clean", seg["source"])
        tgt = seg.get("target_clean", seg["target"])
        status = "Translated" if tgt.strip() else ("Fuzzy" if seg.get("fuzzy") else "Untranslated")
        ws.append([idx + 1, src, tgt if tgt.strip() else "", status])

    wb.save(output_path)


def export_pdf(segments, output_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors

    doc = SimpleDocTemplate(output_path, pagesize=A4,
                            leftMargin=20*mm, rightMargin=20*mm,
                            topMargin=20*mm, bottomMargin=20*mm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title2", parent=styles["Heading1"], fontSize=16, spaceAfter=12)
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8, leading=10)

    elements = []
    elements.append(Paragraph("Translation Export", title_style))
    elements.append(Spacer(1, 6*mm))

    data = [["#", "Source", "Target"]]
    for idx, seg in enumerate(segments):
        src = seg.get("source_clean", seg["source"])
        tgt = seg.get("target_clean", seg["target"])
        data.append([
            str(idx + 1),
            Paragraph(src[:200] if len(src) > 200 else src, cell_style),
            Paragraph(tgt[:200] if len(tgt) > 200 else tgt, cell_style) if tgt.strip() else "",
        ])

    col_widths = [12*mm, 77*mm, 77*mm]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(table)
    doc.build(elements)
