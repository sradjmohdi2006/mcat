import unittest
import os
import tempfile
import json
import shutil
import sys

from cat_tool.project import Project
from cat_tool.file_handler import CATFileHandler
from cat_tool.formats.mqxliif import MqxliffHandler
from cat_tool.formats.sdlxliff import SdlxliffHandler
from cat_tool.formats.text import TextHandler
from cat_tool.formats._text_base import make_segments_from_texts


# ========================== PROJECT TESTS ==========================

class TestProject(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def test_new_project(self):
        p = Project()
        p.new("Test Proj", "en", "es", ["a.txt", "b.txt"])
        self.assertEqual(p.name, "Test Proj")
        self.assertEqual(p.source_lang, "en")
        self.assertEqual(p.target_lang, "es")
        self.assertEqual(len(p.source_files), 2)
        self.assertTrue(p.created_at)
        self.assertTrue(p.modified_at)

    def test_save_and_load(self):
        p = Project()
        p.new("SaveTest", "fr", "de", [])
        proj_path = os.path.join(self.tmp.name, "test.mcatproj")
        p.save(proj_path)
        self.assertTrue(os.path.exists(proj_path))
        loaded = Project.load(proj_path)
        self.assertEqual(loaded.name, "SaveTest")
        self.assertEqual(loaded.source_lang, "fr")
        self.assertEqual(loaded.target_lang, "de")

    def test_save_no_path_raises(self):
        p = Project()
        p.new("NoPath", "en", "es", [])
        with self.assertRaises(ValueError):
            p.save()

    def test_to_dict_roundtrip(self):
        p = Project()
        p.new("DictTest", "en", "es", ["f1.txt"])
        d = p.to_dict()
        self.assertEqual(d["version"], 1)
        self.assertEqual(d["name"], "DictTest")
        p2 = Project.from_dict(d)
        self.assertEqual(p2.name, "DictTest")
        self.assertEqual(p2.source_lang, "en")

    def test_default_dir(self):
        d = Project.default_dir()
        self.assertTrue(os.path.isabs(d))

    def test_recent_projects(self):
        Project.add_recent("/fake/path", "FakeProj")
        recents = Project.get_recent()
        found = any(r.get("path") == "/fake/path" for r in recents)
        self.assertTrue(found)
        self.assertLessEqual(len(recents), 10)

    def test_str(self):
        p = Project()
        p.new("StrTest", "en", "es", [])
        s = str(p)
        self.assertIn("StrTest", s)
        self.assertIn("en", s)

    def test_omegat_save_and_load(self):
        src_file = os.path.join(self.tmp.name, "doc.txt")
        with open(src_file, "w", encoding="utf-8") as f:
            f.write("Hello")
        p = Project()
        p.new("OmegaTTest", "en", "es", [src_file])
        omegat_dir = os.path.join(self.tmp.name, "OmegaTTest")
        p.save_as_omegat(omegat_dir)
        self.assertTrue(os.path.exists(os.path.join(omegat_dir, "omegat.project")))
        self.assertTrue(os.path.exists(os.path.join(omegat_dir, "source", "doc.txt")))
        loaded = Project.load_omegat(omegat_dir)
        self.assertEqual(loaded.name, "OmegaTTest")
        self.assertEqual(loaded.source_lang, "en")
        self.assertEqual(loaded.target_lang, "es")
        self.assertIn("doc.txt", loaded.source_files[0])

    def test_omegat_load_raises_on_missing(self):
        with self.assertRaises(FileNotFoundError):
            Project.load_omegat("nonexistent_dir")

    def test_memoq_save_and_load(self):
        src_file = os.path.join(self.tmp.name, "doc.txt")
        with open(src_file, "w", encoding="utf-8") as f:
            f.write("Hello")
        p = Project()
        p.new("MemoQTest", "en", "es", [src_file])

        handler = CATFileHandler()
        handler.segments = make_segments_from_texts(["Hello"], notes_prefix="")
        handler.segments[0]["target"] = "Hola"

        memoq_dir = os.path.join(self.tmp.name, "memoq_out")
        p.save_as_memoq(memoq_dir, handler)
        self.assertTrue(os.path.exists(os.path.join(memoq_dir, "manifest.xml")))
        self.assertTrue(os.path.exists(os.path.join(memoq_dir, "source", "doc.txt")))
        self.assertTrue(os.path.exists(os.path.join(memoq_dir, "target", "MemoQTest.mqxliff")))

        loaded = Project.load_memoq(memoq_dir)
        self.assertEqual(loaded.name, "MemoQTest")
        self.assertEqual(loaded.source_lang, "en")
        self.assertEqual(loaded.target_lang, "es")
        self.assertGreater(len(loaded.mqxliff_files), 0)

    def test_memoq_load_raises_on_missing(self):
        with self.assertRaises(FileNotFoundError):
            Project.load_memoq("nonexistent_dir")

    def test_load_omegat_from_minimal_xml(self):
        omegat_dir = os.path.join(self.tmp.name, "minimal_omegat")
        os.makedirs(omegat_dir)
        os.makedirs(os.path.join(omegat_dir, "source"))
        et = ET = __import__("xml.etree.ElementTree", fromlist=["ElementTree"])
        root = et.Element("omegat")
        pe = et.SubElement(root, "project")
        et.SubElement(pe, "source_lang").text = "EN-US"
        et.SubElement(pe, "target_lang").text = "FR"
        tree = et.ElementTree(root)
        tree.write(os.path.join(omegat_dir, "omegat.project"), encoding="utf-8", xml_declaration=True)
        loaded = Project.load_omegat(omegat_dir)
        self.assertEqual(loaded.source_lang, "en-us")
        self.assertEqual(loaded.target_lang, "fr")

    def test_load_omegat_missing_project_element(self):
        omegat_dir = os.path.join(self.tmp.name, "bad_omegat")
        os.makedirs(omegat_dir)
        et = ET = __import__("xml.etree.ElementTree", fromlist=["ElementTree"])
        root = et.Element("omegat")
        tree = et.ElementTree(root)
        tree.write(os.path.join(omegat_dir, "omegat.project"), encoding="utf-8", xml_declaration=True)
        with self.assertRaises(ValueError):
            Project.load_omegat(omegat_dir)


# ========================== MQXLIFF TESTS ==========================

class TestMqxliffHandler(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def _make_mqxliff(self, content):
        path = os.path.join(self.tmp.name, "test.mqxliff")
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path

    def test_load_simple(self):
        xml = """<?xml version="1.0"?>
<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2" version="1.2">
  <file original="test" source-language="en" target-language="es">
    <body>
      <trans-unit id="1">
        <source>Hello</source>
        <target>Hola</target>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        path = self._make_mqxliff(xml)
        segs = MqxliffHandler().load(path)
        self.assertEqual(len(segs), 1)
        self.assertEqual(segs[0]["source"], "Hello")
        self.assertEqual(segs[0]["target"], "Hola")
        self.assertTrue(segs[0]["translated"])

    def test_load_empty_target(self):
        xml = """<?xml version="1.0"?>
<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2" version="1.2">
  <file original="test" source-language="en" target-language="es">
    <body>
      <trans-unit id="1">
        <source>Hello</source>
        <target></target>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        path = self._make_mqxliff(xml)
        segs = MqxliffHandler().load(path)
        self.assertEqual(len(segs), 1)
        self.assertEqual(segs[0]["target"], "")
        self.assertFalse(segs[0]["translated"])

    def test_load_missing_body_raises(self):
        xml = """<?xml version="1.0"?>
<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2" version="1.2">
  <file original="test" source-language="en" target-language="es">
  </file>
</xliff>"""
        path = self._make_mqxliff(xml)
        with self.assertRaises(ValueError):
            MqxliffHandler().load(path)

    def test_save_and_reload(self):
        segs = [{
            "index": 0, "source": "Hello", "target": "Hola",
            "source_clean": "Hello", "target_clean": "Hola",
            "all_tags": [], "notes": "",
            "fuzzy": False, "translated": True,
            "locations": "", "store": None, "unit": None,
        }]
        out = os.path.join(self.tmp.name, "out.mqxliff")
        MqxliffHandler().save(segs, out, src_lang="en", tgt_lang="es")
        self.assertTrue(os.path.exists(out))
        reloaded = MqxliffHandler().load(out)
        self.assertEqual(len(reloaded), 1)
        self.assertEqual(reloaded[0]["source"], "Hello")
        self.assertEqual(reloaded[0]["target"], "Hola")

    def test_save_no_target(self):
        segs = [{
            "index": 0, "source": "Hello", "target": "",
            "source_clean": "Hello", "target_clean": "",
            "all_tags": [], "notes": "",
            "fuzzy": False, "translated": False,
            "locations": "", "store": None, "unit": None,
        }]
        out = os.path.join(self.tmp.name, "empty.mqxliff")
        MqxliffHandler().save(segs, out)
        self.assertTrue(os.path.exists(out))
        reloaded = MqxliffHandler().load(out)
        self.assertEqual(reloaded[0]["target"], "")


# ========================== SDLXLIFF TESTS ==========================

class TestSdlxliffHandler(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def _make_sdlxliff(self, content):
        path = os.path.join(self.tmp.name, "test.sdlxliff")
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path

    def test_load_simple(self):
        xml = """<?xml version="1.0"?>
<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2"
       xmlns:sdl="http://sdl.com/FileTypes/SdlXliff/1.0" version="1.2">
  <file original="test" datatype="x-sdlfilterframework2" source-language="en" target-language="es">
    <header/>
    <body>
      <trans-unit id="1">
        <source>Hello</source>
        <seg-source><mrk mtype="seg" mid="1">Hello</mrk></seg-source>
        <target><mrk mtype="seg" mid="1">Hola</mrk></target>
        <sdl:seg-defs><sdl:seg id="1" conf="Translated"/></sdl:seg-defs>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        path = self._make_sdlxliff(xml)
        segs = SdlxliffHandler().load(path)
        self.assertEqual(len(segs), 1)
        self.assertEqual(segs[0]["source"], "Hello")
        self.assertEqual(segs[0]["target"], "Hola")
        self.assertTrue(segs[0]["translated"])

    def test_load_skips_translate_no(self):
        xml = """<?xml version="1.0"?>
<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2"
       xmlns:sdl="http://sdl.com/FileTypes/SdlXliff/1.0" version="1.2">
  <file original="test" datatype="x-sdlfilterframework2" source-language="en" target-language="es">
    <header/>
    <body>
      <trans-unit id="1" translate="no">
        <source>Skip me</source>
      </trans-unit>
      <trans-unit id="2">
        <source>Hello</source>
        <seg-source><mrk mtype="seg" mid="1">Hello</mrk></seg-source>
        <target><mrk mtype="seg" mid="1">Hola</mrk></target>
        <sdl:seg-defs><sdl:seg id="1" conf="Translated"/></sdl:seg-defs>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        path = self._make_sdlxliff(xml)
        segs = SdlxliffHandler().load(path)
        self.assertEqual(len(segs), 1)
        self.assertEqual(segs[0]["source"], "Hello")

    def test_load_skip_no_source(self):
        xml = """<?xml version="1.0"?>
<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2"
       xmlns:sdl="http://sdl.com/FileTypes/SdlXliff/1.0" version="1.2">
  <file original="test" datatype="x-sdlfilterframework2" source-language="en" target-language="es">
    <header/>
    <body>
      <trans-unit id="1">
        <seg-source><mrk mtype="seg" mid="1"></mrk></seg-source>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        path = self._make_sdlxliff(xml)
        segs = SdlxliffHandler().load(path)
        self.assertEqual(len(segs), 0)

    def test_save_and_reload(self):
        segs = [{
            "index": 0, "source": "Hello", "target": "Hola",
            "source_clean": "Hello", "target_clean": "Hola",
            "all_tags": [], "notes": "",
            "fuzzy": False, "translated": True,
            "locations": "", "store": None, "unit": None,
        }]
        out = os.path.join(self.tmp.name, "out.sdlxliff")
        SdlxliffHandler().save(segs, out, src_lang="en", tgt_lang="es")
        self.assertTrue(os.path.exists(out))
        reloaded = SdlxliffHandler().load(out)
        self.assertEqual(len(reloaded), 1)
        self.assertEqual(reloaded[0]["source"], "Hello")
        self.assertEqual(reloaded[0]["target"], "Hola")

    def test_save_without_target(self):
        segs = [{
            "index": 0, "source": "Hello", "target": "",
            "source_clean": "Hello", "target_clean": "",
            "all_tags": [], "notes": "",
            "fuzzy": False, "translated": False,
            "locations": "", "store": None, "unit": None,
        }]
        out = os.path.join(self.tmp.name, "empty.sdlxliff")
        SdlxliffHandler().save(segs, out)
        self.assertTrue(os.path.exists(out))
        reloaded = SdlxliffHandler().load(out)
        self.assertEqual(reloaded[0]["target"], "")

    def test_target_without_mrk(self):
        xml = """<?xml version="1.0"?>
<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2"
       xmlns:sdl="http://sdl.com/FileTypes/SdlXliff/1.0" version="1.2">
  <file original="test" datatype="x-sdlfilterframework2" source-language="en" target-language="es">
    <header/>
    <body>
      <trans-unit id="1">
        <source>Hello</source>
        <seg-source><mrk mtype="seg" mid="1">Hello</mrk></seg-source>
        <target>Hola</target>
        <sdl:seg-defs><sdl:seg id="1" conf="Translated"/></sdl:seg-defs>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        path = self._make_sdlxliff(xml)
        segs = SdlxliffHandler().load(path)
        self.assertEqual(len(segs), 1)
        self.assertEqual(segs[0]["target"], "Hola")


# ========================== GUI SMOKE TEST ==========================

class TestGuiSmoke(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def test_create_mainwindow(self):
        from PyQt5.QtWidgets import QApplication
        app = QApplication.instance()
        existed = app is not None
        if not existed:
            app = QApplication(sys.argv)
        try:
            from cat_tool.ui import MainWindow
            win = MainWindow()
            self.assertIsNotNone(win)
            self.assertEqual(win.windowTitle()[:4], "mcat")
        finally:
            if not existed:
                pass
            app = None

    def test_load_po_file_through_handler(self):
        from PyQt5.QtWidgets import QApplication
        app = QApplication.instance()
        existed = app is not None
        if not existed:
            app = QApplication(sys.argv)
        try:
            from cat_tool.ui import MainWindow
            win = MainWindow()
            po_content = '''msgid "Hello"
msgstr "Hola"
msgid "World"
msgstr "Mundo"
'''
            po_path = os.path.join(self.tmp.name, "test.po")
            with open(po_path, "w", encoding="utf-8") as f:
                f.write(po_content)
            segs = win.file_handler.load_file(po_path)
            win.populate_table(segs)
            self.assertEqual(win.table.rowCount(), 2)
            self.assertEqual(win.table.item(0, 0).text(), "Hello")
        finally:
            if not existed:
                pass
            app = None

    def test_load_po_and_translate(self):
        from PyQt5.QtWidgets import QApplication
        app = QApplication.instance()
        existed = app is not None
        if not existed:
            app = QApplication(sys.argv)
        try:
            from cat_tool.ui import MainWindow
            win = MainWindow()
            po_content = '''msgid "Hello"
msgstr ""
'''
            po_path = os.path.join(self.tmp.name, "test.po")
            with open(po_path, "w", encoding="utf-8") as f:
                f.write(po_content)
            segs = win.file_handler.load_file(po_path)
            win.populate_table(segs)
            win.file_handler.segments[0]["target"] = "Hola"
            win.update_table_row_status(0, win.file_handler.segments[0])
            self.assertEqual(win.file_handler.segments[0]["target"], "Hola")
        finally:
            if not existed:
                pass
            app = None

    def test_qa_dialog_creation(self):
        from PyQt5.QtWidgets import QApplication, QDialog
        app = QApplication.instance()
        existed = app is not None
        if not existed:
            app = QApplication(sys.argv)
        try:
            from cat_tool.qa import QAIssue
            from cat_tool.ui import QADialog
            issues = [
                QAIssue("error", "Missing translation", 0, "Hello", "", "No target"),
            ]
            dlg = QADialog(issues)
            self.assertIsInstance(dlg, QDialog)
        finally:
            if not existed:
                pass
            app = None

    def test_shortcuts_dialog_creation(self):
        from PyQt5.QtWidgets import QApplication, QDialog
        app = QApplication.instance()
        existed = app is not None
        if not existed:
            app = QApplication(sys.argv)
        try:
            from cat_tool.ui import ShortcutsDialog
            dlg = ShortcutsDialog()
            self.assertIsInstance(dlg, QDialog)
        finally:
            if not existed:
                pass
            app = None

    def test_glossary_dialog_creation(self):
        from PyQt5.QtWidgets import QApplication, QDialog
        app = QApplication.instance()
        existed = app is not None
        if not existed:
            app = QApplication(sys.argv)
        try:
            from cat_tool.ui import GlossaryDialog
            dlg = GlossaryDialog()
            self.assertIsInstance(dlg, QDialog)
        finally:
            if not existed:
                pass
            app = None


# ========================== FORMAT HANDLER INTEGRATION ==========================

class TestFileHandlerIntegration(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def _make_po(self, content):
        path = os.path.join(self.tmp.name, "test.po")
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path

    def test_po_load_save_roundtrip(self):
        po = '''msgid "Hello"
msgstr "Hola"
msgid "World"
msgstr ""
'''
        path = self._make_po(po)
        fh = CATFileHandler()
        segs = fh.load_file(path)
        segs[0]["target"] = "¡Hola!"
        segs[1]["target"] = "Mundo"
        out = os.path.join(self.tmp.name, "out.po")
        fh.save_file(out)
        fh2 = CATFileHandler()
        segs2 = fh2.load_file(out)
        self.assertEqual(segs2[0]["target"], "¡Hola!")
        self.assertEqual(segs2[1]["target"], "Mundo")

    def test_txt_load_save_roundtrip(self):
        txt_path = os.path.join(self.tmp.name, "input.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Line one\nLine two\nLine three")
        fh = CATFileHandler()
        segs = fh.load_file(txt_path)
        self.assertEqual(len(segs), 3)
        segs[0]["target"] = "Línea uno"
        segs[1]["target"] = "Línea dos"
        out = os.path.join(self.tmp.name, "out.txt")
        fh.render_translated(out)
        with open(out, encoding="utf-8") as f:
            content = f.read()
        self.assertIn("Línea uno", content)
        self.assertIn("Línea dos", content)

    def test_po_convert_to_po_from_txt(self):
        from cat_tool.file_handler import convert_to_po
        txt_path = os.path.join(self.tmp.name, "source.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Hello\nWorld")
        po_out = os.path.join(self.tmp.name, "converted.po")
        convert_to_po(txt_path, po_out)
        self.assertTrue(os.path.exists(po_out))
        fh = CATFileHandler()
        segs = fh.load_file(po_out)
        self.assertEqual(len(segs), 2)

    def test_factory_xliff_load(self):
        xliff = """<?xml version="1.0" encoding="utf-8"?>
<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2" version="1.2">
  <file source-language="en" target-language="es" datatype="plaintext" original="test.txt">
    <body>
      <trans-unit id="1">
        <source>Hello</source>
        <target>Hola</target>
      </trans-unit>
    </body>
  </file>
</xliff>"""
        path = os.path.join(self.tmp.name, "test.xlf")
        with open(path, "w", encoding="utf-8") as f:
            f.write(xliff)
        fh = CATFileHandler()
        segs = fh.load_file(path)
        self.assertGreater(len(segs), 0)
        self.assertEqual(segs[0]["source"].strip(), "Hello")

    def test_render_no_current_path_raises(self):
        fh = CATFileHandler()
        fh.segments = [{"source": "Hello", "target": "Hola"}]
        with self.assertRaises(ValueError):
            fh.render_translated("out.txt")

    def test_load_mcatdb_and_save_mcatdb_roundtrip(self):
        from cat_tool.mcat_format_manager import MCatProjectManager
        db_path = os.path.join(self.tmp.name, "proj.mcat.db")
        MCatProjectManager.create_project_file(db_path, "en", "es")
        MCatProjectManager.import_extracted_segments(db_path, [
            {"file_id": "f.txt", "source": "Hello", "target": "Hola", "status": "Confirmed", "match_score": 100},
        ])
        fh = CATFileHandler()
        segs = fh.load_file(db_path)
        self.assertEqual(len(segs), 1)
        self.assertEqual(segs[0]["source"], "Hello")
        self.assertEqual(segs[0]["target"], "Hola")
        out_db = os.path.join(self.tmp.name, "copy.mcat.db")
        fh.save_file(out_db)
        self.assertTrue(os.path.exists(out_db))
        fh2 = CATFileHandler()
        segs2 = fh2.load_file(out_db)
        self.assertEqual(segs2[0]["source"], "Hello")


# ========================== REAL DOCUMENT LOAD TESTS ==========================

class TestRealDocumentLoad(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def test_docx_load(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Hello World")
        doc.add_paragraph("Second paragraph")
        docx_path = os.path.join(self.tmp.name, "test.docx")
        doc.save(docx_path)
        from cat_tool.formats.docx import DocxHandler
        segs = DocxHandler().load(docx_path)
        self.assertGreaterEqual(len(segs), 2)
        texts = [s["source"] for s in segs]
        self.assertIn("Hello World", texts)
        self.assertIn("Second paragraph", texts)

    def test_docx_render(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Hello")
        docx_path = os.path.join(self.tmp.name, "test.docx")
        doc.save(docx_path)
        from cat_tool.formats.docx import DocxHandler
        handler = DocxHandler()
        segs = handler.load(docx_path)
        segs[0]["target"] = "Hola"
        out_path = os.path.join(self.tmp.name, "out.docx")
        handler.render(docx_path, segs, out_path)
        self.assertTrue(os.path.exists(out_path))
        doc2 = Document(out_path)
        texts = [p.text for p in doc2.paragraphs]
        self.assertIn("Hola", texts)

    def test_pptx_load(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        tf = txBox.text_frame
        tf.text = "Hello from PowerPoint"
        pptx_path = os.path.join(self.tmp.name, "test.pptx")
        prs.save(pptx_path)
        from cat_tool.formats.pptx import PptxHandler
        segs = PptxHandler().load(pptx_path)
        texts = [s["source"] for s in segs]
        self.assertIn("Hello from PowerPoint", texts)

    def test_pptx_render(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        tf = txBox.text_frame
        tf.text = "Hello"
        pptx_path = os.path.join(self.tmp.name, "test.pptx")
        prs.save(pptx_path)
        from cat_tool.formats.pptx import PptxHandler
        handler = PptxHandler()
        segs = handler.load(pptx_path)
        segs[0]["target"] = "Hola"
        out_path = os.path.join(self.tmp.name, "out.pptx")
        handler.render(pptx_path, segs, out_path)
        self.assertTrue(os.path.exists(out_path))
        prs2 = Presentation(out_path)
        for slide in prs2.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    self.assertIn("Hola", shape.text_frame.text)

    def test_xlsx_load(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Hello"
        ws["B1"] = "World"
        xlsx_path = os.path.join(self.tmp.name, "test.xlsx")
        wb.save(xlsx_path)
        from cat_tool.formats.xlsx import XlsxHandler
        segs = XlsxHandler().load(xlsx_path)
        texts = [s["source"] for s in segs]
        self.assertIn("Hello", texts)
        self.assertIn("World", texts)

    def test_xlsx_render(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Hello"
        xlsx_path = os.path.join(self.tmp.name, "test.xlsx")
        wb.save(xlsx_path)
        from cat_tool.formats.xlsx import XlsxHandler
        handler = XlsxHandler()
        segs = handler.load(xlsx_path)
        segs[0]["target"] = "Hola"
        out_path = os.path.join(self.tmp.name, "out.xlsx")
        handler.render(xlsx_path, segs, out_path)
        self.assertTrue(os.path.exists(out_path))
        from openpyxl import load_workbook
        wb2 = load_workbook(out_path)
        self.assertEqual(wb2.active["A1"].value, "Hola")


if __name__ == "__main__":
    unittest.main()
